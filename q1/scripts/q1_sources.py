"""Official Q1 source audit and reversible text tokenization."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re
import subprocess

from openpyxl import load_workbook


TOKEN_RE = re.compile(r"[A-Za-z]+(?:['’][A-Za-z]+)*|\d+(?:[.,]\d+)*")


def official_words(raw: str) -> list[dict]:
    """Retain original character spans; normalization never changes raw_text."""
    return [{"word": m.group(), "align_word": m.group().replace("’", "'").lower(),
             "char_start": m.start(), "char_end": m.end()}
            for m in TOKEN_RE.finditer(raw)]


def sample_id(video_id: str, clip_id: str) -> str:
    return f"{video_id}$_${clip_id}"


def safe_id(video_id: str, clip_id: str) -> str:
    # Hex encoding is reversible and safe for all filesystem names.
    return f"{video_id.encode().hex()}__{clip_id.encode().hex()}"


def load_official_labels(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    names = [str(x).strip().lower() if x is not None else "" for x in next(rows)]
    required = {"video_id", "clip_id", "text"}
    if not required.issubset(names):
        raise ValueError(f"label sheet missing columns: {sorted(required - set(names))}")
    result = []
    seen = set()
    for row in rows:
        if not any(x is not None for x in row):
            continue
        data = dict(zip(names, row))
        video_id = str(data["video_id"]).strip()
        clip_id = str(data["clip_id"]).strip()
        if clip_id.endswith(".0") and isinstance(data["clip_id"], (int, float)):
            clip_id = clip_id[:-2]
        key = (video_id, clip_id)
        if key in seen:
            raise ValueError(f"duplicate official key: {key}")
        seen.add(key)
        result.append({"video_id": video_id, "clip_id": clip_id,
                       "sample_id": sample_id(*key), "safe_id": safe_id(*key),
                       "raw_text": str(data["text"] or "")})
    wb.close()
    return result


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ffprobe(path: Path, command: str = "ffprobe") -> dict:
    result = subprocess.run([command, "-v", "error", "-show_format", "-show_streams",
                             "-of", "json", str(path)], capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(f"ffprobe failed: {result.stderr.strip()}")
    return json.loads(result.stdout)


def stream_info(probe: dict, kind: str) -> dict | None:
    streams = [s for s in probe.get("streams", []) if s.get("codec_type") == kind]
    return streams[0] if streams else None


def number(x, default=0.0) -> float:
    try:
        return float(x)
    except (ValueError, TypeError):
        return default


def media_summary(probe: dict) -> dict:
    fmt = probe.get("format", {})
    video = stream_info(probe, "video")
    audio = stream_info(probe, "audio")
    container_start = number(fmt.get("start_time"))
    def span(stream):
        if not stream:
            return None
        start = number(stream.get("start_time"), container_start) - container_start
        duration = number(stream.get("duration"), number(fmt.get("duration")))
        return [max(0.0, start), max(0.0, start) + duration]
    fps_str = (video or {}).get("avg_frame_rate", "0/1")
    try:
        n, d = map(float, fps_str.split("/"))
        fps = n / d if d else 0.0
    except ValueError:
        fps = 0.0
    return {"duration_s": number(fmt.get("duration")), "container_start_s": container_start,
            "video_span_s": span(video), "audio_span_s": span(audio),
            "fps": fps, "audio_sr": int((audio or {}).get("sample_rate", 0)),
            "video_stream": video, "audio_stream": audio}


def decode_ok(path: Path, kind: str, command: str = "ffmpeg") -> bool:
    selector = "v:0" if kind == "video" else "a:0"
    result = subprocess.run([command, "-nostdin", "-v", "error", "-i", str(path),
                             "-map", f"0:{selector}", "-f", "null", "-"],
                            capture_output=True, text=True)
    return result.returncode == 0
